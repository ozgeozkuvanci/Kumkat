#!/usr/bin/env python3
"""run_recipes.py için örnek girdi dosyaları üretir (Pillow, python-pptx, ffmpeg, soffice gerekir)."""
import pathlib, subprocess, sys, tempfile

out = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "samples")
out.mkdir(parents=True, exist_ok=True)
lo = pathlib.Path(tempfile.mkdtemp(prefix="lo_")).as_uri()

from PIL import Image, ImageDraw
im = Image.new("RGBA", (800, 600), (0, 0, 0, 0))
ImageDraw.Draw(im).ellipse((100, 50, 700, 550), fill=(255, 140, 0, 255))
im.save(out / "foto.png")

# Photoshop yapısında, RLE sıkıştırmalı birleşik görüntülü PSD
import struct
rgb = im.convert("RGB"); W, H = rgb.size
def packbits(row):
    res, i = bytearray(), 0
    while i < len(row):
        n = min(128, len(row) - i)
        res += bytes([n - 1]) + row[i:i + n]; i += n
    return bytes(res)
rows = [packbits(rgb.getchannel(c).tobytes()[y * W:(y + 1) * W]) for c in "RGB" for y in range(H)]
psd = b"8BPS" + struct.pack(">H6xHIIHH", 1, 3, H, W, 8, 3) + struct.pack(">I", 0) * 3 + struct.pack(">H", 1)
psd += b"".join(struct.pack(">H", len(r)) for r in rows) + b"".join(rows)
(out / "foto.psd").write_bytes(psd)

def run(*cmd):
    subprocess.run(cmd, check=True, cwd=out, capture_output=True)

run("ffmpeg", "-y", "-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30:duration=4",
    "-f", "lavfi", "-i", "sine=frequency=440:duration=4",
    "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", "klip.mov")
run("ffmpeg", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=3", "ses.wav")

(out / "not.txt").write_text("Merhaba Dünya\nİkinci satır ğüşıöç\n", encoding="utf-8")
(out / "tablo.csv").write_text('ad,yaş,şehir\nEnes,33,İstanbul\n"Ali, Veli",28,"Gebze ""Kocaeli"""\n', encoding="utf-8")

from pptx import Presentation
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = "Kumkat"; s.placeholders[1].text = "Test"
p.save(out / "sunum.pptx")

base = ["soffice", "--headless", "--norestore", f"-env:UserInstallation={lo}"]
run(*base, "--convert-to", "docx:MS Word 2007 XML", "not.txt", "--outdir", ".")
run(*base, "--infilter=CSV:44,34,76,1", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "tablo.csv", "--outdir", ".")
run(*base, "--convert-to", "pdf", "sunum.pptx", "--outdir", ".")
print("örnekler:", sorted(f.name for f in out.iterdir()))
